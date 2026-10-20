from io import BytesIO
from typing import cast
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl import Workbook, load_workbook

from app.routers.uploads import extract_csv_metadata, extract_excel_metadata
from app.services.api_exports import excel_response


def test_excel_blank_header_does_not_shift_column_values():
    workbook = Workbook()
    cast(Worksheet, workbook.active).append(["customerCode", None, "segmentCode"])
    cast(Worksheet, workbook.active).append(["C1", "ignored", "vip"])
    content = BytesIO()
    workbook.save(content)
    assert extract_excel_metadata(content.getvalue())[2] == [{"customerCode": "C1", "segmentCode": "vip"}]


def test_exported_user_text_is_not_a_formula():
    response = excel_response(["Name"], [["=HYPERLINK(\"https://example.com\")"]], "test.xlsx")
    workbook = load_workbook(BytesIO(response.body))
    worksheet = workbook.active
    assert worksheet is not None
    assert worksheet["A2"].data_type == "s"
    workbook.close()


def test_csv_header_whitespace_preserves_values():
    assert extract_csv_metadata(b"customerCode , segmentCode\nC1,vip\n")[2] == [{"customerCode": "C1", "segmentCode": "vip"}]
