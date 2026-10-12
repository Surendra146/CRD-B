"""Regression coverage for request contracts exercised by the React UI."""
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Customer, DataUpload, User

pytestmark = pytest.mark.integration


def create_customer(client, account, payload, **extra):
    response = client.post("/api/customers/", headers=account["headers"], json={**payload, **extra})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_customer_sort_and_sales_metrics(client, account, customer_payload):
    first = create_customer(client, account, customer_payload, externalId="A", lifecycle={"totalSpent": 9})
    second = create_customer(client, account, customer_payload, externalId="B", name="Other", lifecycle={"totalSpent": 100})
    result = client.get("/api/customers/", headers=account["headers"], params={"sortBy": "lifecycle.totalSpent", "sortOrder": "desc"})
    assert [row["id"] for row in result.json()["data"]] == [second["id"], first["id"]]
    result = client.put(f"/api/customers/{first['id']}", headers=account["headers"], json={
        "customerCreatedDate": "2026-10-04T00:00:00.000Z", "purchases": [{"amount": 123, "status": "completed"}]})
    assert result.status_code == 200, result.text
    assert result.json()["data"]["lifecycle"]["totalSpent"] == 123
    assert result.json()["data"]["lifecycle"]["totalPurchases"] == 1
    assert client.put(f"/api/customers/{first['id']}", headers=account["headers"], json={"customerCreatedDate": "bad"}).status_code == 400


def test_segment_import_validates_saves_and_exports(client, account, customer_payload):
    customer = create_customer(client, account, customer_payload)
    segment = client.post("/api/segments/", headers=account["headers"], json={"name": "VIP", "code": "vip"})
    assert segment.status_code == 200
    template = client.get("/api/segments/import/template", headers=account["headers"])
    assert template.content[:2] == b"PK"
    workbook = load_workbook(BytesIO(template.content))
    assert list(workbook.active.values)[0] == ("customerCode", "segmentCode")
    workbook.close()
    content = b"customerCode,segmentCode\nCUST-001,vip\nmissing,vip\n"
    response = client.post("/api/segments/import/validate", headers=account["headers"], files={"file": ("mappings.csv", content, "text/csv")})
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert (data["totalRows"], data["validCount"], data["errorCount"]) == (2, 1, 1)
    saved = client.post("/api/segments/import/save", headers=account["headers"], json={"validRows": data["validRows"]})
    assert saved.json()["data"]["savedRows"] == 1
    assert client.get(f"/api/customers/{customer['id']}", headers=account["headers"]).json()["data"]["lifecycle"]["segment"] == "vip"
    exported = client.post("/api/segments/import/export-errors", headers=account["headers"], json={"errors": data["errors"]})
    assert exported.content[:2] == b"PK"


def test_segment_save_cannot_use_another_tenants_codes(client, account_factory, customer_payload):
    owner, other = account_factory(), account_factory()
    create_customer(client, other, customer_payload)
    assert client.post("/api/segments/", headers=other["headers"], json={"name": "Other", "code": "other"}).status_code == 200
    response = client.post("/api/segments/import/save", headers=owner["headers"], json={"validRows": [{"customerCode": "CUST-001", "segmentCode": "other"}]})
    assert response.status_code == 400


def test_report_contract_filters_and_binary_downloads(client, account_factory, db_session):
    owner, other = account_factory(), account_factory()
    now = datetime.now(UTC).replace(tzinfo=None)
    for account, upload_type in ((owner, "customer_details"), (owner, "customer_sales"), (other, "customer_details")):
        user = db_session.get(User, account["user"]["id"])
        db_session.add(DataUpload(tenant_id=user.tenant_id, organization_id=user.organization_id, uploaded_by=user.id,
                                  type=upload_type, file={"originalName": "example.csv"}, stats={"totalRows": 3}, created_at=now))
    db_session.commit()
    options = client.get("/api/uploads/history/report/options", headers=owner["headers"]).json()["data"]
    assert options["users"] == [{"id": owner["user"]["id"], "name": "Test owner"}]
    assert len(options["uploadTypes"]) == 2
    params = {"reportType": "uploaded_excel_type_wise", "uploadType": "customer_sales", "fromDate": now.date().isoformat(), "toDate": now.date().isoformat()}
    response = client.get("/api/uploads/history/report", headers=owner["headers"], params=params)
    assert response.status_code == 200, response.text
    assert response.json()["data"]["total"] == 1
    assert len(response.json()["data"]["rows"]) == 1
    excel = client.get("/api/uploads/history/report/export/excel", headers=owner["headers"], params=params)
    assert excel.content[:2] == b"PK"
    workbook = load_workbook(BytesIO(excel.content))
    assert len(list(workbook.active.values)) == 2
    workbook.close()
    pdf = client.get("/api/uploads/history/report/export/pdf", headers=owner["headers"], params=params)
    assert pdf.content.startswith(b"%PDF-")
    assert client.get("/api/uploads/history", headers=owner["headers"], params={"limit": 1}).json()["data"]
    assert client.get("/api/uploads/history/report", headers=owner["headers"], params={**params, "fromDate": "bad"}).status_code == 400
    assert client.get("/api/uploads/history/report", headers=owner["headers"], params={"userId": other["user"]["id"]}).json()["data"]["total"] == 0


def test_campaign_filters_and_schedule_contract(client, account, customer_payload):
    create_customer(client, account, customer_payload, lifecycle={"daysSinceLastPurchase": 10, "totalSpent": 20, "totalPurchases": 1})
    create_customer(client, account, customer_payload, externalId="SECOND", name="Another", lifecycle={"daysSinceLastPurchase": 45, "totalSpent": 500, "totalPurchases": 8})
    payload = {"name": "Filtered", "type": "one_time", "audience": {"type": "segment", "filters": {"minDaysSinceLastPurchase": 30, "minTotalSpent": 100, "minOrders": 2}},
               "schedule": {"scheduledAt": (datetime.now(UTC) + timedelta(days=2)).isoformat()}}
    response = client.post("/api/campaigns/", headers=account["headers"], json=payload)
    assert response.status_code == 200, response.text
    assert response.json()["data"]["audience"]["estimatedReach"] == 1
    assert response.json()["data"]["schedule"]["scheduledAt"] == payload["schedule"]["scheduledAt"]
    assert client.post("/api/campaigns/", headers=account["headers"], json={**payload, "schedule": {"startDate": "2030-01-01"}}).status_code == 400


def test_logout_clears_cookie_even_if_token_expired(client):
    client.cookies.set("token", "expired")
    response = client.post("/api/auth/logout")
    assert response.status_code == 200
    assert 'token=""' in response.headers["set-cookie"]


def test_cors_exposes_retry_and_download_headers(client):
    response = client.get("/health", headers={"Origin": "https://frontend.example.com"})
    assert response.headers["access-control-expose-headers"] == "Retry-After, Content-Disposition"


def test_readiness_logs_failed_dependency_without_credentials(client, monkeypatch, caplog):
    import app.main as main

    class BrokenEngine:
        def connect(self):
            raise RuntimeError("private connection credential")

    monkeypatch.setattr(main, "engine", BrokenEngine())
    response = client.get("/ready")
    assert response.status_code == 503
    assert response.json() == {"success": False, "status": "unavailable"}
    assert "runtime_database (RuntimeError)" in caplog.text
    assert "private connection credential" not in caplog.text
