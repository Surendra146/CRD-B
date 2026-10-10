import asyncio
import csv
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO, StringIO
from typing import Any

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Query, UploadFile
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.database.connection import SessionLocal, ControlSessionLocal, get_db
from app.database.tenancy import bind_scope
from app.models import Customer, DataUpload, DataUploadRow, User
from app.redis.connection import get_progress_cache
from app.schemas.upload import ColumnMappingRequest, SuggestMappingsRequest
from app.services.security import ensure_tenant_access, require_organization, tenant_id_for_user
from app.socket.progress import emit_upload_progress
from app.utils.helpers import apply_customer_purchase_metrics, day_bounds, json_ready, model_to_dict, normalize_phone


PREVIEW_LIMIT = 100
STAGING_INSERT_CHUNK_SIZE = 1000
SAVE_CHUNK_SIZE = 500


TARGET_FIELDS_BY_UPLOAD_TYPE = {
    "customer_details": [
        {"value": "demographics.location.locationName", "label": "Location Name"},
        {"value": "customerFirstName", "label": "Customer First Name"},
        {"value": "customerLastName", "label": "Customer Last Name"},
        {"value": "phone", "label": "Phone Number"},
        {"value": "customerCreatedDate", "label": "Customer Created Date"},
        {
            "value": "demographics.customerType",
            "label": "Customer Type",
            "defaultValue": "Regular",
        },
        {"value": "demographics.gender", "label": "Gender"},
        {"value": "demographics.location.country", "label": "Country"},
        {"value": "demographics.location.state", "label": "State"},
        {"value": "demographics.location.city", "label": "City"},
        {"value": "address", "label": "Address"},
    ],
    "customer_sales": [
        {"value": "demographics.location.locationName", "label": "Location Name"},
        {"value": "customerFirstName", "label": "Customer First Name"},
        {"value": "customerLastName", "label": "Customer Last Name"},
        {"value": "phone", "label": "Phone Number"},
        {"value": "bill_number", "label": "Bill Number"},
        {"value": "bill_date", "label": "Bill Date"},
        {"value": "counter_number", "label": "Counter Number"},
        {"value": "unit_price", "label": "Unit Price"},
        {"value": "quantity", "label": "Quantity"},
        {
            "value": "total_price",
            "label": "Total Price",
            "derivedFrom": ["unit_price", "quantity"],
        },
    ],
}

REQUIRED_FIELDS_BY_UPLOAD_TYPE = {
    "customer_details": ["name", "phone"],
    "customer_sales": ["name", "phone", "bill_number", "bill_date", "unit_price", "quantity"],
}

TARGET_FIELD_ALIASES = {
    "demographics.location.locationName": ["location", "location name", "store", "store name", "branch"],
    "customerFirstName": ["first name", "customer first name", "firstname", "customer firstname"],
    "customerLastName": ["last name", "customer last name", "lastname", "customer lastname"],
    "phone": ["phone", "phone number", "mobile", "mobile number", "contact", "contact number"],
    "customerCreatedDate": ["created date", "customer created date", "registration date", "signup date"],
    "demographics.customerType": ["customer type", "type"],
    "demographics.gender": ["gender", "sex"],
    "demographics.location.country": ["country"],
    "demographics.location.state": ["state"],
    "demographics.location.city": ["city"],
    "address": ["address", "customer address"],
    "bill_number": ["bill number", "bill no", "bill#", "invoice number", "invoice no", "receipt number"],
    "bill_date": ["bill date", "invoice date", "receipt date", "purchase date", "date"],
    "counter_number": ["counter number", "counter no", "pos", "pos no", "terminal"],
    "unit_price": ["unit price", "per item price", "item price", "rate", "price"],
    "quantity": ["quantity", "qty", "item quantity"],
    "total_price": ["total price", "per item total price", "line total", "amount", "total"],
}


def normalize_column_name(value: str) -> str:
    return " ".join(str(value or "").replace("_", " ").replace("-", " ").lower().split())


def target_fields_for_type(upload_type: str | None) -> list[dict]:
    return TARGET_FIELDS_BY_UPLOAD_TYPE.get(upload_type or "", TARGET_FIELDS_BY_UPLOAD_TYPE["customer_details"])


def suggest_target_field(column: str, upload_type: str | None) -> str:
    normalized_column = normalize_column_name(column)
    fields = target_fields_for_type(upload_type)
    valid_targets = {field["value"] for field in fields}

    for field in fields:
        if normalized_column == normalize_column_name(field["label"]):
            return field["value"]

    for target_field, aliases in TARGET_FIELD_ALIASES.items():
        if target_field not in valid_targets:
            continue
        if normalized_column in {normalize_column_name(alias) for alias in aliases}:
            return target_field

    return ""


def normalize_column_mapping(mappings: list, upload_type: str | None) -> list:
    normalized = [mapping for mapping in mappings if isinstance(mapping, dict)]
    mapped_targets = {
        mapping.get("targetField")
        for mapping in normalized
        if isinstance(mapping.get("targetField"), str) and mapping.get("targetField")
    }

    if upload_type == "customer_details" and "demographics.customerType" not in mapped_targets:
        normalized.append(
            {
                "sourceColumn": "__default_customer_type",
                "targetField": "demographics.customerType",
                "transformation": "default",
                "defaultValue": "Regular",
            }
        )

    if (
        upload_type == "customer_sales"
        and "total_price" not in mapped_targets
        and {"unit_price", "quantity"}.issubset(mapped_targets)
    ):
        normalized.append(
            {
                "sourceColumn": "__derived_total_price",
                "targetField": "total_price",
                "transformation": "multiply",
                "derivedFrom": ["unit_price", "quantity"],
            }
        )

    return normalized


def get_upload_metadata(upload: DataUpload, columns: list[str] | None = None) -> dict:
    data = model_to_dict(upload)
    stats = data.get("stats") or {}
    data["uploadId"] = data["id"]
    data["columns"] = columns if columns is not None else data.get("columns") or []
    data["totalRows"] = stats.get("totalRows") or stats.get("total_rows") or 0
    data["previewLimit"] = PREVIEW_LIMIT
    return data


def stage_upload_rows(db: Session, upload: DataUpload, rows: list[dict]) -> None:
    staged_rows: list[DataUploadRow] = []
    for index, row in enumerate(rows, start=2):
        staged_rows.append(
            DataUploadRow(
                upload_id=upload.id,
                tenant_id=upload.tenant_id,
                organization_id=upload.organization_id,
                row_number=index,
                status="staged",
                raw_data=row,
                mapped_data={},
                errors=[],
                warnings=[],
            )
        )
        if len(staged_rows) >= STAGING_INSERT_CHUNK_SIZE:
            db.bulk_save_objects(staged_rows)
            staged_rows = []

    if staged_rows:
        db.bulk_save_objects(staged_rows)


def get_staged_rows(db: Session, upload: DataUpload, status: str | None = None):
    statement = (
        select(DataUploadRow)
        .where(
            DataUploadRow.upload_id == upload.id,
            DataUploadRow.tenant_id == upload.tenant_id,
            DataUploadRow.organization_id == upload.organization_id,
        )
        .order_by(DataUploadRow.row_number)
    )
    if status:
        statement = statement.where(DataUploadRow.status == status)
    return db.scalars(statement)


def count_staged_rows(db: Session, upload: DataUpload, status: str | None = None) -> int:
    statement = select(func.count(DataUploadRow.id)).where(
        DataUploadRow.upload_id == upload.id,
        DataUploadRow.tenant_id == upload.tenant_id,
        DataUploadRow.organization_id == upload.organization_id,
    )
    if status:
        statement = statement.where(DataUploadRow.status == status)
    return int(db.scalar(statement) or 0)


def preview_rows(db: Session, upload: DataUpload, status: str, limit: int = PREVIEW_LIMIT) -> list[dict]:
    rows = db.scalars(
        select(DataUploadRow)
        .where(
            DataUploadRow.upload_id == upload.id,
            DataUploadRow.tenant_id == upload.tenant_id,
            DataUploadRow.organization_id == upload.organization_id,
            DataUploadRow.status == status,
        )
        .order_by(DataUploadRow.row_number)
        .limit(limit)
    ).all()

    return [
        {
            "row": row.row_number,
            "data": row.mapped_data if status in {"valid", "saved"} else row.raw_data,
            "message": "; ".join(row.errors or []),
            "warnings": row.warnings or [],
        }
        for row in rows
    ]


def warning_preview_rows(db: Session, upload: DataUpload, limit: int = PREVIEW_LIMIT) -> list[dict]:
    warnings = []
    for row in get_staged_rows(db, upload):
        if not row.warnings:
            continue
        warnings.append(
            {
                "row": row.row_number,
                "data": row.mapped_data or row.raw_data,
                "message": "; ".join(row.warnings or []),
                "warnings": row.warnings or [],
            }
        )
        if len(warnings) >= limit:
            break
    return warnings


def upload_response(upload: DataUpload, db: Session | None = None) -> dict:
    data = model_to_dict(upload)
    data["uploadId"] = data["id"]
    if db is not None:
        data["warningPreview"] = warning_preview_rows(db, upload)
    return data


def progress_response(upload: DataUpload, db: Session | None = None, include_previews: bool = False) -> dict:
    if include_previews:
        return upload_response(upload, db)
    data = model_to_dict(upload)
    data["uploadId"] = data["id"]
    return data


def publish_upload_progress(upload: DataUpload, db: Session | None = None, include_previews: bool = False) -> None:
    payload = progress_response(upload, db, include_previews)
    asyncio.run(emit_upload_progress(upload.id, payload))


def phone_counts_for_upload(db: Session, upload: DataUpload) -> Counter:
    counts: Counter = Counter()
    for row in get_staged_rows(db, upload):
        mapped = map_import_row(row.raw_data or {}, upload.column_mapping or [])
        phone = normalize_phone(mapped.get("phone"))
        if phone:
            counts[phone] += 1
    return counts


def existing_phones_for_upload(db: Session, upload: DataUpload, phones: set[str]) -> set[str]:
    existing: set[str] = set()
    phone_list = list(phones)
    for index in range(0, len(phone_list), STAGING_INSERT_CHUNK_SIZE):
        batch = phone_list[index : index + STAGING_INSERT_CHUNK_SIZE]
        existing.update(
            phone
            for phone in db.scalars(
                select(Customer.phone).where(
                    Customer.organization_id == upload.organization_id,
                    Customer.tenant_id == upload.tenant_id,
                    Customer.phone.in_(batch),
                )
            ).all()
            if phone is not None
        )
    return existing


def parse_number(value: Any) -> float:
    try:
        return float(str(value or "0").replace(",", "").strip() or 0)
    except ValueError:
        return 0


def parse_date(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    bounds = day_bounds(value)
    return bounds[0] if bounds else None


def set_nested_value(payload: dict, path: str, value: Any) -> None:
    current = payload
    parts = path.split(".")
    for part in parts[:-1]:
        current = current.setdefault(part, {})
    current[parts[-1]] = value


def row_value(row: dict, mapping: dict, mapped: dict | None = None) -> Any:
    transformation = mapping.get("transformation")
    if transformation == "default":
        return mapping.get("defaultValue")
    if transformation == "multiply":
        left, right = mapping.get("derivedFrom") or [None, None]
        source = mapped or row
        return parse_number(source.get(left)) * parse_number(source.get(right))
    return row.get(mapping.get("sourceColumn"))


def map_import_row(row: dict, mappings: dict | list) -> dict:
    mapped: dict[str, Any] = {}
    for mapping in mappings if isinstance(mappings, list) else []:
        if not isinstance(mapping, dict):
            continue
        target = mapping.get("targetField")
        if not target:
            continue
        value = row_value(row, mapping, mapped)
        if value is None or value == "":
            continue
        mapped[target] = value
        if "." in target:
            set_nested_value(mapped, target, value)

    if "customerFirstName" in mapped or "customerLastName" in mapped:
        mapped["name"] = " ".join(
            str(mapped.get(part) or "").strip()
            for part in ["customerFirstName", "customerLastName"]
            if str(mapped.get(part) or "").strip()
        )
    mapped.setdefault("demographics", {})
    mapped["demographics"].setdefault("customerType", "Regular")
    return mapped


def build_customer_payload(mapped: dict, user: User, upload_type: str) -> dict:
    customer_date = parse_date(mapped.get("customerCreatedDate")) or datetime.now(timezone.utc).replace(tzinfo=None)
    return {
        "organization_id": user.organization_id,
        "tenant_id": tenant_id_for_user(user),
        "tenant_code": user.tenant_code,
        "external_id": normalize_phone(mapped.get("phone")) or None,
        "name": mapped.get("name") or "Unknown Customer",
        "email": mapped.get("email"),
        "phone": normalize_phone(mapped.get("phone")),
        "address": mapped.get("address"),
        "customer_created_date": customer_date,
        "demographics": mapped.get("demographics") or {"customerType": "Regular"},
        "lifecycle": mapped.get("lifecycle") or {"status": "new"},
        "preferences": {"preferredChannel": "whatsapp", "marketingOptIn": False, "language": "en"},
        "source": {"type": "import", "importedAt": datetime.now(timezone.utc).isoformat()},
        "module_tags": [upload_type],
    }


def get_or_create_import_customer(db: Session, mapped: dict, user: User, upload_type: str) -> tuple[Customer, bool]:
    phone = normalize_phone(mapped.get("phone"))
    customer = None
    if phone:
        customer = db.scalar(
            select(Customer).where(
                Customer.organization_id == user.organization_id,
                Customer.tenant_id == tenant_id_for_user(user),
                Customer.phone == phone,
            )
        )

    created = False
    if not customer:
        customer = Customer(**build_customer_payload(mapped, user, upload_type))
        db.add(customer)
        created = True
    else:
        tags = set(customer.module_tags or [])
        tags.add(upload_type)
        customer.module_tags = list(tags)
        if mapped.get("name"):
            customer.name = mapped["name"]
        if mapped.get("address"):
            customer.address = mapped["address"]
        demographics = customer.demographics or {}
        incoming_demographics = mapped.get("demographics") or {}
        demographics.update(incoming_demographics)
        customer.demographics = demographics

    return customer, created


def build_purchase(mapped: dict) -> dict:
    unit_price = parse_number(mapped.get("unit_price"))
    quantity = parse_number(mapped.get("quantity"))
    total_price = parse_number(mapped.get("total_price")) or unit_price * quantity
    bill_date = parse_date(mapped.get("bill_date"))
    location = (mapped.get("demographics") or {}).get("location") or {}
    return {
        "orderId": mapped.get("bill_number"),
        "billNumber": mapped.get("bill_number"),
        "posNo": mapped.get("counter_number"),
        "counterNumber": mapped.get("counter_number"),
        "date": bill_date.isoformat() if bill_date else None,
        "amount": total_price,
        "locationName": location.get("locationName"),
        "items": [
            {
                "unit_price": unit_price,
                "quantity": quantity,
                "total_price": total_price,
                "price": unit_price,
            }
        ],
        "status": "completed",
    }


def extract_csv_metadata(content: bytes) -> tuple[list[str], int, list[dict]]:
    text = content.decode("utf-8-sig", errors="ignore")
    reader = csv.DictReader(StringIO(text))
    columns = [str(column).strip() for column in (reader.fieldnames or []) if str(column).strip()]
    rows = [
        {column: json_ready(row.get(column)) for column in columns}
        for row in reader
    ]
    return columns, len(rows), rows


def extract_excel_metadata(content: bytes) -> tuple[list[str], int, list[dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(400, "XLSX uploads require openpyxl to read columns") from exc

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        workbook.close()
        return [], 0, []
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or []
    columns = [str(column).strip() for column in header if column is not None and str(column).strip()]
    parsed_rows = []
    for values in rows:
        parsed_rows.append(
            {
                column: json_ready(values[index] if index < len(values) else None)
                for index, column in enumerate(columns)
            }
        )
    workbook.close()
    return columns, len(parsed_rows), parsed_rows


def extract_upload_metadata(filename: str | None, content: bytes) -> tuple[list[str], int, list[dict]]:
    extension = (filename or "").lower().rsplit(".", 1)[-1]
    if extension == "csv":
        return extract_csv_metadata(content)
    if extension == "xlsx":
        return extract_excel_metadata(content)
    if extension == "xls":
        raise HTTPException(400, "XLS uploads are not supported yet. Please upload CSV or XLSX.")
    return [], 0, []


async def upload_file(
    type: str = Form("customer_details"),
    file: UploadFile = File(...),
    user: User = Depends(require_organization),
    db: Session = Depends(get_db),
):
    content = await file.read()
    columns, total_rows, rows = extract_upload_metadata(file.filename, content)
    raw_preview = rows[:PREVIEW_LIMIT]
    upload = DataUpload(
        tenant_id=tenant_id_for_user(user),
        organization_id=user.organization_id,
        tenant_code=user.tenant_code,
        uploaded_by=user.id,
        file={
            "originalName": file.filename,
            "storedName": file.filename,
            "mimeType": file.content_type,
            "size": len(content),
        },
        type=type,
        status="staged",
        stats={
            "totalRows": total_rows,
            "processedRows": 0,
            "successRows": 0,
            "errorRows": 0,
            "warningRows": 0,
            "previewLimit": PREVIEW_LIMIT,
            "staging": "data_upload_rows",
        },
        valid_preview=[{"columns": columns, "rows": raw_preview}],
        error_preview=[],
        errors=[],
    )
    db.add(upload)
    db.commit()
    db.refresh(upload)
    stage_upload_rows(db, upload, rows)
    db.commit()
    return {"success": True, "data": get_upload_metadata(upload, columns)}


def suggest_mappings(payload: SuggestMappingsRequest):
    payload_data = payload.to_payload()
    columns = payload_data.get("columns") or []
    upload_type = payload_data.get("type")
    return {
        "success": True,
        "data": [
            {
                "sourceColumn": column,
                "targetField": suggest_target_field(column, upload_type),
                "transformation": "none",
            }
            for column in columns
        ],
    }


def get_target_fields(type: str = Query("customer_details")):
    return {
        "success": True,
        "data": target_fields_for_type(type),
    }


def get_upload_history(user: User = Depends(require_organization), db: Session = Depends(get_db)):
    rows = db.scalars(
        select(DataUpload)
        .where(DataUpload.tenant_id == tenant_id_for_user(user), DataUpload.organization_id == user.organization_id)
        .order_by(DataUpload.created_at.desc())
    ).all()
    return {"success": True, "data": [model_to_dict(row) for row in rows]}


def get_upload_report_options():
    return {"success": True, "data": {"statuses": ["pending", "processing", "completed", "failed", "partial"]}}


def get_upload_report(user: User = Depends(require_organization), db: Session = Depends(get_db)):
    return get_upload_history(user, db)


def export_upload_report_excel():
    return {"success": True, "message": "Excel export is ready to implement with openpyxl"}


def export_upload_report_pdf():
    return {"success": True, "message": "PDF export is ready to implement"}


def set_column_mapping(upload_id: int, payload: ColumnMappingRequest, user: User = Depends(require_organization), db: Session = Depends(get_db)):
    payload_data = payload.to_payload()
    upload = db.get(DataUpload, upload_id)
    upload = ensure_tenant_access(upload, user, "Upload")
    requested_mapping = payload_data.get("columnMapping") or payload_data.get("mappings") or []
    upload.column_mapping = normalize_column_mapping(requested_mapping, upload.type)
    db.commit()
    return {"success": True, "data": upload_response(upload, db)}


def validate_upload_job(upload_id: int) -> None:
    with ControlSessionLocal() as discovery:
        owner = discovery.get(DataUpload, upload_id)
        if not owner:
            return
        tenant_id, organization_id = owner.tenant_id, owner.organization_id
    db = SessionLocal()
    bind_scope(db, tenant_id, organization_id)
    try:
        upload = db.get(DataUpload, upload_id)
        if not upload:
            return

        phone_counts = phone_counts_for_upload(db, upload)
        existing_phones = existing_phones_for_upload(db, upload, set(phone_counts.keys()))
        processed_rows = 0
        valid_rows = 0
        error_rows = 0
        warning_rows = 0
        total_rows = count_staged_rows(db, upload)

        for row in get_staged_rows(db, upload):
            mapped = map_import_row(row.raw_data or {}, upload.column_mapping or [])
            errors = []
            warnings = []

            for field in REQUIRED_FIELDS_BY_UPLOAD_TYPE.get(upload.type, REQUIRED_FIELDS_BY_UPLOAD_TYPE["customer_details"]):
                if not str(mapped.get(field) or "").strip():
                    errors.append(f"Missing required field: {field}")

            phone = normalize_phone(mapped.get("phone"))
            if phone and phone_counts.get(phone, 0) > 1:
                warnings.append("Duplicate phone number found in this file")
            if phone and phone in existing_phones:
                warnings.append("Existing customer will be updated on save")

            if upload.type == "customer_sales":
                unit_price = parse_number(mapped.get("unit_price"))
                quantity = parse_number(mapped.get("quantity"))
                if unit_price <= 0:
                    errors.append("Unit price must be greater than 0")
                if quantity <= 0:
                    errors.append("Quantity must be greater than 0")
                if not parse_date(mapped.get("bill_date")):
                    errors.append("Bill date is invalid")

            row.mapped_data = mapped
            row.errors = errors
            row.warnings = warnings
            row.status = "invalid" if errors else "valid"

            processed_rows += 1
            if errors:
                error_rows += 1
            else:
                valid_rows += 1
            if warnings:
                warning_rows += 1

            if processed_rows % SAVE_CHUNK_SIZE == 0:
                stats = upload.stats or {}
                upload.stats = {
                    **stats,
                    "totalRows": total_rows,
                    "processedRows": processed_rows,
                    "successRows": valid_rows,
                    "validRows": valid_rows,
                    "errorRows": error_rows,
                    "invalidRows": error_rows,
                    "warningRows": warning_rows,
                    "skippedRows": 0,
                    "stagedRows": total_rows,
                }
                db.commit()
                publish_upload_progress(upload)

        db.flush()
        stats = upload.stats or {}

        upload.status = "validated_with_errors" if error_rows else "validated"
        upload.valid_preview = preview_rows(db, upload, "valid")
        upload.error_preview = preview_rows(db, upload, "invalid")
        upload.errors = upload.error_preview
        upload.stats = {
            **stats,
            "totalRows": total_rows,
            "processedRows": processed_rows,
            "successRows": valid_rows,
            "validRows": valid_rows,
            "errorRows": error_rows,
            "invalidRows": error_rows,
            "warningRows": warning_rows,
            "skippedRows": 0,
            "newCustomers": max(valid_rows - len(existing_phones), 0) if upload.type == "customer_details" else 0,
            "updatedCustomers": 0,
            "newTransactions": valid_rows if upload.type == "customer_sales" else 0,
            "previewLimit": PREVIEW_LIMIT,
            "stagedRows": total_rows,
        }
        db.commit()
        publish_upload_progress(upload, db, include_previews=True)
    except Exception as exc:
        db.rollback()
        upload = db.get(DataUpload, upload_id)
        if upload:
            stats = upload.stats or {}
            upload.status = "failed"
            upload.errors = [{"row": None, "message": str(exc), "data": {}}]
            upload.stats = {**stats, "errorRows": stats.get("errorRows", 0), "failed": True}
            db.commit()
            publish_upload_progress(upload, db, include_previews=True)
    finally:
        db.close()


def process_upload(
    upload_id: int,
    background_tasks: BackgroundTasks,
    user: User = Depends(require_organization),
    db: Session = Depends(get_db),
):
    upload = db.get(DataUpload, upload_id)
    upload = ensure_tenant_access(upload, user, "Upload")
    if not upload.column_mapping:
        raise HTTPException(400, "Please save column mapping before validation")
    if upload.status == "completed":
        raise HTTPException(400, "This upload has already been saved")
    if upload.status == "validating":
        return {"success": True, "data": upload_response(upload, db)}

    upload.status = "validating"
    stats = upload.stats or {}
    upload.stats = {
        **stats,
        "processedRows": 0,
        "successRows": 0,
        "validRows": 0,
        "errorRows": 0,
        "invalidRows": 0,
        "warningRows": 0,
    }
    db.commit()
    db.refresh(upload)
    publish_upload_progress(upload)
    background_tasks.add_task(validate_upload_job, upload.id)
    return {"success": True, "data": upload_response(upload, db)}


def save_upload_job(upload_id: int, user_id: int) -> None:
    with ControlSessionLocal() as discovery:
        owner = discovery.get(DataUpload, upload_id)
        if not owner:
            return
        tenant_id, organization_id = owner.tenant_id, owner.organization_id
    db = SessionLocal()
    bind_scope(db, tenant_id, organization_id)
    try:
        upload = db.get(DataUpload, upload_id)
        user = db.get(User, user_id)
        if not upload or not user:
            return
        if upload.tenant_id != tenant_id_for_user(user) or upload.organization_id != user.organization_id:
            return

        saved_rows = 0
        new_customers = 0
        new_transactions = 0
        customer_cache: dict[str, Customer] = {}
        valid_row_count = count_staged_rows(db, upload, "valid")

        for item in get_staged_rows(db, upload, "valid"):
            mapped = item.mapped_data or {}
            phone = normalize_phone(mapped.get("phone"))
            cached_customer = customer_cache.get(phone) if phone else None
            if cached_customer:
                customer = cached_customer
                created = False
            else:
                customer, created = get_or_create_import_customer(db, mapped, user, upload.type)
                if phone:
                    customer_cache[phone] = customer
            if created:
                new_customers += 1

            if upload.type == "customer_sales":
                purchase = build_purchase(mapped)
                customer.purchases = [*(customer.purchases or []), purchase]
                apply_customer_purchase_metrics(customer)
                new_transactions += 1

            item.status = "saved"
            saved_rows += 1
            if saved_rows % SAVE_CHUNK_SIZE == 0:
                stats = upload.stats or {}
                upload.stats = {
                    **stats,
                    "savedRows": saved_rows,
                    "processedRows": saved_rows,
                    "totalRows": valid_row_count,
                    "newCustomers": new_customers,
                    "newTransactions": new_transactions,
                    "updatedCustomers": max(saved_rows - new_customers, 0),
                }
                db.commit()
                publish_upload_progress(upload)

        db.flush()
        stats = upload.stats or {}
        upload.status = "completed"
        upload.valid_preview = preview_rows(db, upload, "saved")
        upload.stats = {
            **stats,
            "savedRows": saved_rows,
            "processedRows": saved_rows,
            "totalRows": stats.get("stagedRows", stats.get("totalRows", valid_row_count)),
            "newCustomers": stats.get("newCustomers", new_customers) if upload.type == "customer_details" else new_customers,
            "newTransactions": new_transactions if upload.type == "customer_sales" else stats.get("newTransactions", 0),
            "updatedCustomers": max(saved_rows - new_customers, 0),
        }
        upload.saved_at = datetime.now(timezone.utc).replace(tzinfo=None)
        db.commit()
        publish_upload_progress(upload, db, include_previews=True)
    except Exception as exc:
        db.rollback()
        upload = db.get(DataUpload, upload_id)
        if upload:
            stats = upload.stats or {}
            upload.status = "failed"
            upload.errors = [{"row": None, "message": str(exc), "data": {}}]
            upload.stats = {**stats, "failed": True}
            db.commit()
            publish_upload_progress(upload, db, include_previews=True)
    finally:
        db.close()


def confirm_save_upload(
    upload_id: int,
    background_tasks: BackgroundTasks,
    user: User = Depends(require_organization),
    db: Session = Depends(get_db),
):
    upload = db.get(DataUpload, upload_id)
    upload = ensure_tenant_access(upload, user, "Upload")
    if upload.status == "completed":
        return {"success": True, "data": upload_response(upload, db)}
    if upload.status == "saving":
        return {"success": True, "data": upload_response(upload, db)}
    if upload.status not in {"validated", "validated_with_errors"}:
        raise HTTPException(400, "Please validate this upload before saving")

    valid_row_count = count_staged_rows(db, upload, "valid")
    if valid_row_count == 0:
        raise HTTPException(400, "There are no valid rows to save")

    upload.status = "saving"
    stats = upload.stats or {}
    upload.stats = {
        **stats,
        "savedRows": 0,
        "processedRows": 0,
        "totalRows": valid_row_count,
    }
    db.commit()
    db.refresh(upload)
    publish_upload_progress(upload)
    background_tasks.add_task(save_upload_job, upload.id, user.id)
    return {"success": True, "data": upload_response(upload, db)}


def export_upload_errors(upload_id: int, user: User = Depends(require_organization), db: Session = Depends(get_db)):
    upload = db.get(DataUpload, upload_id)
    upload = ensure_tenant_access(upload, user, "Upload")
    return {"success": True, "data": upload.errors or []}


def get_upload_status(upload_id: int, user: User = Depends(require_organization), db: Session = Depends(get_db)):
    upload = db.get(DataUpload, upload_id)
    upload = ensure_tenant_access(upload, user, "Upload")
    if upload.status in {"validated", "validated_with_errors", "completed", "failed", "partial"}:
        return {"success": True, "data": upload_response(upload, db)}
    cached = get_progress_cache(upload_id)
    if (
        cached
        and cached.get("organizationId") == user.organization_id
        and cached.get("tenantId") == tenant_id_for_user(user)
    ):
        return {"success": True, "data": cached}
    return {"success": True, "data": upload_response(upload, db)}


router = APIRouter(dependencies=[Depends(require_organization)])

router.post("/")(upload_file)
router.post("/suggest-mappings")(suggest_mappings)
router.get("/target-fields")(get_target_fields)
router.get("/history")(get_upload_history)
router.get("/history/report/options")(get_upload_report_options)
router.get("/history/report")(get_upload_report)
router.get("/history/report/export/excel")(export_upload_report_excel)
router.get("/history/report/export/pdf")(export_upload_report_pdf)
router.put("/{upload_id}/mapping")(set_column_mapping)
router.post("/{upload_id}/process")(process_upload)
router.post("/{upload_id}/confirm-save")(confirm_save_upload)
router.get("/{upload_id}/errors/export")(export_upload_errors)
router.get("/{upload_id}/status")(get_upload_status)




